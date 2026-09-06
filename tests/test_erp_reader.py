from datetime import datetime

import pytest
from openpyxl import Workbook

from ingest.erp_reader import read_trs_declarations


HEADER = ['Début Equipe', 'Num Equipe', 'Réf. Machine', 'Réf OF', 'Type OF',
          'Qté Pieces Fab.', 'Qté Pieces Bonnes', 'Total Rebuts', 'Nb Cycles',
          'Nb Emp Réelle', 'Taux Rebuts']


def workbook(tmp_path, rows, header=HEADER, duplicate_sheet=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'RESULTAT_EQUIPE'
    sheet['A1'] = 'Résultats par équipe'
    for index, value in enumerate(header, 1):
        sheet.cell(11, index, value)
    for row in rows:
        sheet.append(row)
    if duplicate_sheet:
        book.copy_worksheet(sheet)
    path = tmp_path / 'teams.xlsx'
    book.save(path)
    return path


def row(**changes):
    values = [datetime(2026, 9, 5, 5), 1, '606', '000012', 'Production', 800, 760, 40, 100, 8, .05]
    for key, value in changes.items():
        values[int(key)] = value
    return values


def test_same_order_preserves_two_team_declarations(tmp_path):
    path = workbook(tmp_path, [row(), row(**{'0': datetime(2026, 9, 5, 13), '1': 2, '5': 960, '6': 944, '7': 16, '8': 120})])
    result = read_trs_declarations(path, source_timezone='Europe/Paris')
    assert result.header_row == 11
    assert not result.issues
    assert [d.order_ref for d in result.declarations] == ['000012', '000012']
    assert [d.shift_number for d in result.declarations] == [1, 2]
    assert [d.produced_parts for d in result.declarations] == [800, 960]
    assert [d.cycle_count for d in result.declarations] == [100, 120]
    assert result.declarations[0].shift_started_at.hour == 3
    assert result.declarations[0].order_target_quantity is None
    assert not result.declarations[0].provided_order_fields


@pytest.mark.parametrize('ratio', [.05, 5, '5%'])
def test_ratios_are_normalized_once(tmp_path, ratio):
    result = read_trs_declarations(workbook(tmp_path, [row(**{'10': ratio})]), source_timezone='UTC')
    assert result.declarations[0].declared_scrap_rate == .05


def test_unknown_quality_zero_production_team_four_and_raw_columns(tmp_path):
    result = read_trs_declarations(workbook(tmp_path, [row(**{'1': 4, '5': 0, '6': None, '7': None}) + ['opaque']], HEADER + ['Inconnue']), source_timezone='UTC')
    declaration = result.declarations[0]
    assert declaration.good_parts is None and declaration.scrap_parts is None
    assert declaration.produced_parts == 0 and declaration.shift_number == 4
    assert declaration.raw_data['Inconnue'] == 'opaque'


@pytest.mark.parametrize('value,code', [(None, 'required'), ('25/10/2026 02:30', 'ambiguous_local_time'), ('29/03/2026 02:30', 'nonexistent_local_time')])
def test_invalid_dates_are_public_issues(tmp_path, value, code):
    result = read_trs_declarations(workbook(tmp_path, [row(**{'0': value})]), source_timezone='Europe/Paris')
    assert code in [issue.code for issue in result.issues]
    assert not result.declarations


def test_text_and_excel_dates(tmp_path):
    for value in ['05/09/2026 05:00', 46270 + 5 / 24]:
        result = read_trs_declarations(workbook(tmp_path, [row(**{'0': value})]), source_timezone='UTC')
        assert not result.issues
        assert result.declarations[0].shift_started_at.hour == 5


def test_collision_and_sheet_selection_are_blocking(tmp_path):
    result = read_trs_declarations(workbook(tmp_path, [row(), row()]), source_timezone='UTC')
    assert 'duplicate_declaration_key' in [issue.code for issue in result.issues]
    path = workbook(tmp_path, [row()], duplicate_sheet=True)
    assert read_trs_declarations(path, source_timezone='UTC').issues[0].code == 'sheet_selection_required'
    assert not read_trs_declarations(path, source_timezone='UTC', sheet_name='RESULTAT_EQUIPE').issues


def test_header_absent(tmp_path):
    result = read_trs_declarations(workbook(tmp_path, [], ['Other']), source_timezone='UTC')
    assert result.issues[0].code == 'header_not_found'


def test_signed_erp_totals_are_preserved_with_nonblocking_warnings(tmp_path):
    result = read_trs_declarations(workbook(tmp_path, [row(**{'5': 0, '6': -16, '7': 16}) + [-.08], [''] * 12], HEADER + ['T.R.S.']), source_timezone='UTC')
    assert len(result.declarations) == 1
    assert result.declarations[0].good_parts == -16
    assert result.declarations[0].declared_trs == -.08
    assert {issue.field for issue in result.issues} == {'good_parts', 'declared_trs'}
    assert all(issue.severity == 'warning' for issue in result.issues)
    assert result.declarations[0].warnings


def test_matching_final_summary_is_control_total_not_declaration(tmp_path):
    footer = [None, None, None, None, None, 1600, 1520, 80, 200, None, None]
    result = read_trs_declarations(workbook(tmp_path, [row(), row(**{'1': 2}), footer]), source_timezone='UTC')
    assert len(result.declarations) == 2
    assert len(result.issues) == 1 and result.issues[0].severity == 'info'
    assert result.issues[0].code == 'verified_control_total'
    footer[6] = 1500
    invalid = read_trs_declarations(workbook(tmp_path, [row(), row(**{'1': 2}), footer]), source_timezone='UTC')
    assert any(issue.severity == 'error' for issue in invalid.issues)



def test_blank_optional_order_cells_do_not_erase_known_context(tmp_path):
    result = read_trs_declarations(workbook(tmp_path, [row() + [10000, 'open'], row(**{'1': 2}) + [None, None]], HEADER + ['Qté Cible OF', 'Statut OF']), source_timezone='UTC')
    assert result.declarations[0].provided_order_fields == {'order_target_quantity', 'order_status'}
    assert result.declarations[1].provided_order_fields == set()


def test_excel_percent_format_is_already_a_fraction_even_above_one(tmp_path):
    from openpyxl import load_workbook
    path = workbook(tmp_path, [row(**{'10': 1.25})])
    book = load_workbook(path)
    book.active.cell(12, 11).number_format = '0.00%'
    book.save(path)
    result = read_trs_declarations(path, source_timezone='UTC')
    assert result.declarations[0].declared_scrap_rate == 1.25
    assert result.issues[0].severity == 'warning'
