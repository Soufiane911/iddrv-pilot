"""Read real TRS workbooks without turning shifts or cycles into orders."""
from datetime import date, datetime, time, timezone
import math
from pathlib import Path
import re
import unicodedata
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .erp_models import ERPDeclaration, ERPReadResult, ERPRowIssue, json_value


def normalize(value):
    return re.sub(r'[^a-z0-9]', '', unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode().lower())


ALIASES = {
    'shift_started_at': ['Début Equipe', 'started_at'],
    'shift_number': ['Num Equipe', 'shift_number'],
    'machine_ref': ['Réf. Machine', 'machine_erp_ref'],
    'order_ref': ['Réf OF', 'production_order_id'],
    'order_type': ['Type OF', 'order_type'],
    'tool_ref': ['Réf. outil', 'tool_ref'],
    'product_ref': ['Réf. produit', 'product_ref'],
    'produced_parts': ['Qté Pieces Fab.', 'produced_parts'],
    'good_parts': ['Qté Pieces Bonnes', 'good_parts'],
    'scrap_parts': ['Total Rebuts', 'scrap_count'],
    'cycle_count': ['Nb Cycles', 'cycle_count'],
    'cavities_actual': ['Nb Emp Réelle', 'cavities_actual'],
    'declared_scrap_rate': ['Taux Rebuts', 'declared_scrap_rate'],
    'declared_trs': ['T.R.S.', 'expected_trs'],
    'available_hours': ['Tps Disponible (h)', 'planned_runtime_h'],
    'total_stop_hours': ['Total Arrêts (h)', 'Tps Arrêt Total (h)', 'total_stop_hours'],
    'running_hours': ['Tps Fct Brut (h)', 'running_hours'],
    'opening_hours': ['Tps Ouverture (h)', 'opening_hours'],
    'cycle_time_s': ['Cycle Moyen', 'theoretical_cycle_time_s'],
    'production_started_at': ['Début Production', 'production_started_at'],
    'production_ended_at': ['Fin Equipe', 'ended_at', 'Fin Production', 'production_ended_at'],
    'order_target_quantity': ['Qté Cible OF', 'order_target_quantity', 'target_quantity'],
    'order_status': ['Statut OF', 'order_status'],
    'order_status_effective_at': ['Date Statut OF', 'order_status_effective_at'],
    'order_status_reason': ['Motif Statut OF', 'order_status_reason'],
    'quality_lot_ref': ['Réf Lot Qualité', 'quality_lot_ref'],
    'coverage_complete': ['Historique OF Complet', 'coverage_complete'],
}
MAPPING = {normalize(alias): field for field, aliases in ALIASES.items() for alias in [field, *aliases]}
REQUIRED = {'shift_started_at', 'shift_number', 'machine_ref', 'order_ref'}
ORDER_FIELDS = {'order_target_quantity', 'order_status', 'order_status_effective_at', 'order_status_reason', 'quality_lot_ref', 'coverage_complete'}
INTEGER_FIELDS = {'shift_number', 'produced_parts', 'good_parts', 'scrap_parts', 'cycle_count', 'cavities_actual', 'order_target_quantity'}
FLOAT_FIELDS = {'available_hours', 'total_stop_hours', 'running_hours', 'opening_hours', 'cycle_time_s'}
DATE_FIELDS = {'shift_started_at', 'production_started_at', 'production_ended_at', 'order_status_effective_at'}


def local_datetime(value, source_timezone, epoch=None):
    if isinstance(value, (int, float)):
        value = from_excel(value, epoch=epoch) if epoch else from_excel(value)
    if isinstance(value, date) and not isinstance(value, datetime):
        value = datetime.combine(value, time())
    if not isinstance(value, datetime):
        text = str(value).strip()
        try:
            value = datetime.fromisoformat(text.replace('Z', '+00:00'))
        except ValueError:
            for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y'):
                try:
                    value = datetime.strptime(text, fmt)
                    break
                except ValueError:
                    continue
            else:
                raise ValueError('invalid_date') from None
    if value.tzinfo is None:
        zone = ZoneInfo(source_timezone)
        first, second = value.replace(tzinfo=zone, fold=0), value.replace(tzinfo=zone, fold=1)
        if first.astimezone(timezone.utc).astimezone(zone).replace(tzinfo=None) != value:
            raise ValueError('nonexistent_local_time')
        if first.utcoffset() != second.utcoffset():
            raise ValueError('ambiguous_local_time')
        value = first
    return value.astimezone(timezone.utc)


def identifier(cell):
    value = cell.value
    if isinstance(value, (int, float)) and float(value).is_integer():
        # Excel may store a numeric identifier with a zero-padding display format.
        fmt = cell.number_format
        return str(int(value)).zfill(len(fmt)) if re.fullmatch('0+', fmt) else str(int(value))
    return str(value).strip()


def numeric(value, integer=False, ratio=False, signed=False, ratio_fraction=False):
    text = str(value).strip().replace('\u00a0', '').replace(' ', '').replace(',', '.')
    percentage = text.endswith('%')
    result = float(text.removesuffix('%'))
    if not math.isfinite(result) or (result < 0 and not signed and not ratio):
        raise ValueError('invalid_quantity')
    if ratio:
        result = result / 100 if percentage or (result > 1 and not ratio_fraction) else result
    if integer and not result.is_integer():
        raise ValueError('integer_required')
    return int(result) if integer else result


def read_trs_declarations(path: Path, *, source_timezone: str, sheet_name: str | None = None) -> ERPReadResult:
    ZoneInfo(source_timezone)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        candidates = []
        for sheet in book.worksheets:
            for row in sheet.iter_rows(max_row=30):
                mapped = [MAPPING.get(normalize(cell.value)) for cell in row]
                if REQUIRED.issubset(mapped):
                    candidates.append((sheet, row[0].row, mapped))
                    break
        sheets = [candidate[0].title for candidate in candidates]
        if sheet_name:
            candidates = [candidate for candidate in candidates if candidate[0].title == sheet_name]
        if not candidates or len(candidates) != 1:
            code = 'sheet_selection_required' if len(candidates) > 1 else 'header_not_found'
            return ERPReadResult([], [ERPRowIssue(0, code, None, 'Choisissez une feuille TRS reconnue.' if candidates else 'En-tête TRS introuvable.')], '', 0, sheets)
        sheet, header_row, fields = candidates[0]
        headers = [str(cell.value or f'column_{index}') for index, cell in enumerate(next(sheet.iter_rows(min_row=header_row, max_row=header_row)), 1)]
        declarations, issues, keys = [], [], set()
        mapped_fields = [field for field in fields if field]
        if len(mapped_fields) != len(set(mapped_fields)):
            return ERPReadResult([], [ERPRowIssue(header_row, 'duplicate_column', None, 'Plusieurs colonnes désignent le même champ.')], sheet.title, header_row, sheets)
        def blank(value):
            return value is None or (isinstance(value, str) and not value.strip())
        source_rows = [(number, cells) for number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1)
                       if any(not blank(cell.value) for cell in cells)]
        for row_number, cells in source_rows:
            raw_fields = {field: cell.value for field, cell in zip(fields, cells) if field}
            if row_number == source_rows[-1][0] and declarations and all(blank(raw_fields.get(field)) for field in REQUIRED | {'tool_ref', 'product_ref'}):
                totals = {field: value for field, value in raw_fields.items() if field in {'produced_parts', 'good_parts', 'scrap_parts', 'cycle_count'} and not blank(value)}
                try:
                    valid_total = len(totals) >= 2 and all(
                        all(getattr(declaration, field) is not None for declaration in declarations)
                        and numeric(value, integer=True, signed=True) == sum(getattr(declaration, field) for declaration in declarations)
                        for field, value in totals.items())
                except (ValueError, OverflowError):
                    valid_total = False
                if valid_total:
                    issues.append(ERPRowIssue(row_number, 'verified_control_total', None, 'Total final vérifié ; exclu des déclarations.', 'info'))
                    continue
            values, bad = {}, False
            provided = set()
            warnings = []
            for field, cell in zip(fields, cells):
                if not field:
                    continue
                value = cell.value
                if value is None or (isinstance(value, str) and not value.strip()):
                    if field in REQUIRED:
                        issues.append(ERPRowIssue(row_number, 'required', field, 'Champ obligatoire absent.'))
                        bad = True
                    continue
                try:
                    if field in DATE_FIELDS:
                        value = local_datetime(value, source_timezone, book.epoch)
                    elif field in INTEGER_FIELDS:
                        value = numeric(value, integer=True, signed=field == "good_parts")
                        if field == 'shift_number' and not 1 <= value <= 32767:
                            raise ValueError('invalid_shift_number')
                    elif field in FLOAT_FIELDS:
                        value = numeric(value)
                    elif field in {'declared_trs', 'declared_scrap_rate'}:
                        value = numeric(value, ratio=True, ratio_fraction="%" in cell.number_format)
                    elif field == 'coverage_complete':
                        if str(value).lower() not in {'true', 'false', '1', '0', 'oui', 'non'}:
                            raise ValueError('invalid_boolean')
                        value = str(value).lower() in {'true', '1', 'oui'}
                    else:
                        value = identifier(cell)
                        limit = 50 if field in {'machine_ref', 'order_ref', 'order_type'} else 100
                        if field not in {'order_status_reason'} and len(value) > limit:
                            raise ValueError('identifier_too_long')
                    if (field == 'good_parts' and value < 0) or (field in {'declared_trs', 'declared_scrap_rate'} and not 0 <= value <= 1):
                        warnings.append(field)
                        issues.append(ERPRowIssue(row_number, 'declared_value_to_verify', field, 'Valeur ERP conservée ; résultat à vérifier.', 'warning'))
                    if field in ORDER_FIELDS:
                        provided.add(field)
                    values[field] = value
                except (ValueError, OverflowError):
                    code = 'invalid_value'
                    try:
                        if field in DATE_FIELDS:
                            local_datetime(cell.value, source_timezone, book.epoch)
                    except ValueError as error:
                        code = str(error)
                    issues.append(ERPRowIssue(row_number, code, field, 'Valeur invalide ou heure locale à vérifier.'))
                    bad = True
            if bad:
                continue
            if values.get('production_ended_at'):
                values.setdefault('production_started_at', values['shift_started_at'])
                if values['production_ended_at'] <= values['production_started_at']:
                    issues.append(ERPRowIssue(row_number, 'invalid_period', 'production_ended_at', 'La fin doit suivre le début.'))
                    continue
                values['bounds_origin'] = 'source'
            declaration = ERPDeclaration(**values, raw_data=json_value(dict(zip(headers, [cell.value for cell in cells]))), source_row=row_number, sheet_name=sheet.title, provided_order_fields=provided, warnings=warnings)
            key = declaration.key(0, 0) + declaration.machine_ref
            if key in keys:
                issues.append(ERPRowIssue(row_number, 'duplicate_declaration_key', None, 'Deux lignes partagent la même identité ; corrigez le fichier.'))
            keys.add(key)
            declarations.append(declaration)
        return ERPReadResult(declarations, issues, sheet.title, header_row, sheets)
    finally:
        book.close()
